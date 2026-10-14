#!/usr/bin/env python3
"""The layout description the decks in docs/ are written against, and the two renderers it drives.

A deck is a list of slides, each a list of placed items — a box, a circle, a screenshot, a run of paragraphs — in inches
on a 13.333 x 7.5 canvas. `save_pptx` writes the PowerPoint; `save_preview` writes an HTML replica at the same
coordinates, because this host has no LibreOffice and a deck nobody has looked at is a deck nobody should send.

Copied from srv6-core/docs/decklib.py and extended for homa-lab: lines, native tables and charts, speaker notes, and the
Homa palette. Used by docs/build_deck.py."""
import html as htmlmod
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from lxml import etree
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

W, H = 13.333, 7.5                      # LAYOUT_WIDE

# homa-lab: deep navy for title and divider slides, Homa blue and TCP orange (the PDF's chart colours) as the two
# protagonists, white cards on a cool white ground.
INK, DEEP, TEAL = "0B1A2E", "15294A", "2A78D6"
GREEN, AMBER = "15803D", "EB6834"
GROUND, CARD, LINE = "F5F7FA", "FFFFFF", "DCE2EA"
BODY, MUTED, PALE = "1F2937", "5B6676", "C3D3EA"
HEAD_FONT, BODY_FONT, MONO_FONT = "Cambria", "Calibri", "Courier New"


def para(item, txt, size=None, color=None, bold=False, italic=False, space_after=6):
    """A paragraph inside a text item; size and colour default to the item's."""
    item.setdefault("paras", []).append({"txt": txt, "size": size or item["size"], "color": color or item["color"],
                                         "bold": bold, "italic": italic, "space_after": space_after})


def hexc(c): return RGBColor.from_string(c)


class Deck:
    """One deck. Every helper takes the slide it places into and returns what the caller needs to place the next thing."""

    def __init__(self, shots):
        self.slides, self.shots = [], Path(shots)

    # ---- primitives ------------------------------------------------------------------------------------------
    def slide(self, bg=GROUND, notes=""):
        self.slides.append({"bg": bg, "items": [], "notes": notes}); return self.slides[-1]

    def line(self, s, x1, y1, x2, y2, color=TEAL, width=2):
        s["items"].append({"k": "line", "x": min(x1, x2), "y": min(y1, y2), "w": abs(x2 - x1), "h": abs(y2 - y1),
                           "x1": x1, "y1": y1, "x2": x2, "y2": y2, "color": color, "lw": width})

    def table(self, s, x, y, w, rows, col_w=None, size=13, row_h=0.42, head_fill=DEEP, bold_first_col=False):
        """rows[0] is the header. col_w: relative widths."""
        s["items"].append({"k": "table", "x": x, "y": y, "w": w, "h": row_h * len(rows), "rows": rows, "col_w": col_w,
                           "size": size, "row_h": row_h, "head_fill": head_fill, "bold_first": bold_first_col})

    def chart(self, s, x, y, w, h, kind, cats, series, colors, title="", log=False, unit="", size=12, number_format='0.0#'):
        """A native bar chart. kind: 'bar' (horizontal) or 'col'. series: [(name, [values])]."""
        s["items"].append({"k": "chart", "x": x, "y": y, "w": w, "h": h, "kind": kind, "cats": cats, "series": series,
                           "colors": colors, "title": title, "log": log, "unit": unit, "size": size, "fmt": number_format})

    def text(self, s, x, y, w, h, size=16, color=BODY, bold=False, font=BODY_FONT, align="l", anchor="t",
             spacing=1.15, italic=False):
        s["items"].append({"k": "text", "x": x, "y": y, "w": w, "h": h, "size": size, "color": color, "bold": bold,
                           "font": font, "align": align, "anchor": anchor, "spacing": spacing, "italic": italic})
        return s["items"][-1]

    def box(self, s, x, y, w, h, fill=CARD, line=LINE, radius=0.06, width=1):
        s["items"].append({"k": "box", "x": x, "y": y, "w": w, "h": h, "fill": fill, "line": line,
                           "radius": radius, "lw": width})

    def circle(self, s, x, y, d, fill=TEAL):
        s["items"].append({"k": "circle", "x": x, "y": y, "w": d, "h": d, "fill": fill})

    def size_of(self, name):
        return Image.open(self.shots / f"{name}.png").size

    def shot(self, s, name, x, y, w=None, h=None, frame=True):
        """Place a screenshot, scaled to fit the given width or height without distorting it."""
        iw, ih = self.size_of(name)
        ratio = iw / ih
        if w and not h: h = w / ratio
        elif h and not w: w = h * ratio
        elif w and h:
            if w / h > ratio: w = h * ratio
            else: h = w / ratio
        if frame: self.box(s, x - 0.06, y - 0.06, w + 0.12, h + 0.12, fill=CARD, line=LINE, radius=0.04)
        s["items"].append({"k": "img", "src": str(self.shots / f"{name}.png"), "x": x, "y": y, "w": w, "h": h})
        return w, h

    def caption(self, s, txt, x, y, w, size=11):
        it = self.text(s, x, y, w, 0.35, size=size, color=MUTED)
        para(it, txt, space_after=0)

    def side_note(self, s, x, items, y=1.85, w=None, title=None, h=4.5):
        w = w or (W - x - 0.75)
        self.box(s, x, y, w, h)
        it = self.text(s, x + 0.28, y + 0.25, w - 0.56, h - 0.5, size=12.5, color=BODY, spacing=1.3)
        if title: para(it, title, bold=True, size=14, color=INK, space_after=10)
        for i, t in enumerate(items):
            para(it, t, space_after=0 if i == len(items) - 1 else 9)

    # ---- whole slides ----------------------------------------------------------------------------------------
    def head(self, s, title, sub, tx=0.75, size=30):
        h = self.text(s, tx, 0.5 if size >= 30 else 0.48, W - tx - 0.75, 0.8, size=size, color=INK, bold=True, font=HEAD_FONT)
        para(h, title, space_after=0)
        sb = self.text(s, tx, 1.28 if size >= 30 else 1.15, W - tx - 0.75, 0.45, size=13.5, color=MUTED)
        para(sb, sub, space_after=0)

    def shot_slide(self, title, sub, img, note, img_w=8.2, img_h=4.7, badge=None):
        """A screenshot on the left with room for a card on the right. `badge` numbers a step inside one workflow."""
        s = self.slide()
        tx = 0.75
        if badge:
            self.circle(s, 0.75, 0.5, 0.62)
            num = self.text(s, 0.75, 0.62, 0.62, 0.45, size=17, color="FFFFFF", bold=True, align="c", font=HEAD_FONT)
            para(num, str(badge), space_after=0)
            tx = 1.55
        self.head(s, title, sub, tx=tx, size=28)
        w, hh = self.shot(s, img, 0.75, 1.85, w=img_w, h=img_h)
        self.caption(s, note, 0.75, 1.85 + hh + 0.16, w)
        return s, 0.75 + w + 0.45, hh

    def wide_slide(self, title, sub, img, note, cols, img_h=3.5):
        """For a picture much wider than it is tall: full width, with the commentary in columns underneath (or one line)."""
        s = self.slide()
        self.head(s, title, sub)
        iw, ih = self.size_of(img)                        # a picture narrower than the slide is centred, not left-hung
        ix = 0.75 + max(0.0, (11.83 - min(11.83, img_h * iw / ih)) / 2)
        w, hh = self.shot(s, img, ix, 1.9, w=11.83 - 2 * (ix - 0.75), h=img_h)
        self.caption(s, note, ix, 1.9 + hh + 0.14, w)
        y = 1.9 + hh + 0.62
        if isinstance(cols, str):                         # a picture worth the whole slide gets one line under it
            it = self.text(s, 0.75, y, 11.83, 7.3 - y, size=13, color=BODY, spacing=1.3)
            para(it, cols, space_after=0)
            return s
        cw = (11.83 - 0.4 * (len(cols) - 1)) / len(cols)
        bh = 6.95 - y                                     # the cards run to the bottom margin, so nothing dead is left
        x = 0.75
        for t, b in cols:
            self.box(s, x, y, cw, bh)
            ht = self.text(s, x + 0.28, y + 0.28, cw - 0.56, 0.45, size=14, color=INK, bold=True)
            para(ht, t, space_after=0)
            bt = self.text(s, x + 0.28, y + 0.78, cw - 0.56, bh - 1.06, size=12.5, color=BODY, spacing=1.35)
            para(bt, b, space_after=0)
            x += cw + 0.4
        return s

    def two_up_slide(self, title, sub, left, right, lcap, rcap, notes, band_h=3.4):
        s = self.slide()
        self.head(s, title, sub)

        def fit(name, w, h):
            iw, ih = self.size_of(name)
            r = iw / ih
            return (h * r, h) if w / h > r else (w, w / r)

        lw, lh = fit(left, 5.85, band_h); rw, rh = fit(right, 5.6, band_h)
        band = max(lh, rh)
        self.shot(s, left, 0.75, 1.95 + (band - lh) / 2, w=lw, h=lh)
        self.caption(s, lcap, 0.75, 1.95 + band + 0.14, lw)
        self.shot(s, right, 6.95, 1.95 + (band - rh) / 2, w=rw, h=rh)
        self.caption(s, rcap, 6.95, 1.95 + band + 0.14, rw)
        y = 1.95 + band + 0.62
        it = self.text(s, 0.75, y, 11.8, 7.2 - y, size=12.5, color=BODY, spacing=1.3)
        for i, t in enumerate(notes):
            para(it, t, space_after=0 if i == len(notes) - 1 else 7)
        return s

    def stack_slide(self, title, sub, imgs, caps, notes, y=1.95, gap=0.62):
        """Two wide strips one above the other — for a before/after where the pictures are the comparison."""
        s = self.slide()
        self.head(s, title, sub)
        for img, cap in zip(imgs, caps):
            w, hh = self.shot(s, img, 0.75, y, w=11.83)
            self.caption(s, cap, 0.75, y + hh + 0.13, w)
            y += hh + gap
        it = self.text(s, 0.75, y + 0.05, 11.83, 7.25 - y, size=12.5, color=BODY, spacing=1.3)
        for i, t in enumerate(notes):
            para(it, t, space_after=0 if i == len(notes) - 1 else 8)
        return s

    def part_slide(self, kicker, title, lines):
        """A divider between the halves of a deck."""
        s = self.slide(DEEP)
        k = self.text(s, 1.1, 2.35, 10.5, 0.45, size=13, color=TEAL, bold=True)
        para(k, kicker, space_after=0)
        t = self.text(s, 1.1, 2.85, 10.8, 1.1, size=40, color="FFFFFF", bold=True, font=HEAD_FONT, spacing=1.0)
        para(t, title, space_after=0)
        b = self.text(s, 1.1, 4.2, 10.8, 1.6, size=15, color=PALE, spacing=1.4)
        for i, ln in enumerate(lines):
            para(b, ln, space_after=0 if i == len(lines) - 1 else 8)
        return s

    def columns_slide(self, title, sub, cols, bg=INK, y=2.15, h=2.75, foot=None):
        """Three or four cards across a dark slide — the opening statement, or the closing one."""
        s = self.slide(bg)
        ht = self.text(s, 0.9, 0.9, 11.5, 0.9, size=32, color="FFFFFF", bold=True, font=HEAD_FONT)
        para(ht, title, space_after=0)
        if sub:
            sb = self.text(s, 0.9, 1.72, 11.5, 0.5, size=14.5, color=PALE)
            para(sb, sub, space_after=0)
        cw = (11.52 - 0.2 * (len(cols) - 1)) / len(cols)
        x = 0.9
        for t, b in cols:
            self.box(s, x, y, cw, h, fill=DEEP, line=DEEP)
            tt = self.text(s, x + 0.3, y + 0.27, cw - 0.6, 0.6, size=16, color="FFFFFF", bold=True, font=HEAD_FONT)
            para(tt, t, space_after=0)
            bt = self.text(s, x + 0.3, y + 0.95, cw - 0.6, h - 1.2, size=12, color=PALE, spacing=1.3)
            para(bt, b, space_after=0)
            x += cw + 0.2
        if foot:
            f = self.text(s, 0.9, y + h + 0.6, 11.5, 1.1, size=13, color="7E97A1", spacing=1.35)
            para(f, foot, space_after=0)
        return s

    # ---- renderers -------------------------------------------------------------------------------------------
    def save_pptx(self, out):
        prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
        blank = prs.slide_layouts[6]
        for sp in self.slides:
            sl = prs.slides.add_slide(blank)
            bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(W), Inches(H))
            bg.fill.solid(); bg.fill.fore_color.rgb = hexc(sp["bg"]); bg.line.fill.background(); bg.shadow.inherit = False
            for it in sp["items"]:
                if it["k"] == "box":
                    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if it["radius"] else MSO_SHAPE.RECTANGLE,
                                             Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]))
                    if it["radius"]: sh.adjustments[0] = min(0.5, it["radius"] * 2 / max(it["w"], it["h"]))
                    sh.fill.solid(); sh.fill.fore_color.rgb = hexc(it["fill"])
                    sh.line.color.rgb = hexc(it["line"]); sh.line.width = Pt(it["lw"]); sh.shadow.inherit = False
                elif it["k"] == "circle":
                    sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]))
                    sh.fill.solid(); sh.fill.fore_color.rgb = hexc(it["fill"]); sh.line.fill.background(); sh.shadow.inherit = False
                elif it["k"] == "img":
                    sl.shapes.add_picture(it["src"], Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]))
                elif it["k"] == "line":
                    ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(it["x1"]), Inches(it["y1"]), Inches(it["x2"]), Inches(it["y2"]))
                    ln.line.color.rgb = hexc(it["color"]); ln.line.width = Pt(it["lw"])
                elif it["k"] == "table":
                    rows = it["rows"]; nc = len(rows[0])
                    gt = sl.shapes.add_table(len(rows), nc, Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"])).table
                    cw = it["col_w"] or [1] * nc
                    for j, c in enumerate(cw): gt.columns[j].width = Inches(it["w"] * c / sum(cw))
                    for i, r in enumerate(rows):
                        gt.rows[i].height = Inches(it["row_h"])
                        for j, v in enumerate(r):
                            cell = gt.cell(i, j); cell.text = ""
                            cell.margin_left = cell.margin_right = Inches(0.1); cell.margin_top = cell.margin_bottom = Inches(0.04)
                            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                            cell.fill.solid(); cell.fill.fore_color.rgb = hexc(it["head_fill"] if i == 0 else (CARD if i % 2 else "EEF2F7"))
                            run = cell.text_frame.paragraphs[0].add_run(); run.text = str(v)
                            f = run.font; f.name = BODY_FONT; f.size = Pt(it["size"]); f.bold = i == 0 or (j == 0 and it["bold_first"])
                            f.color.rgb = hexc("FFFFFF" if i == 0 else BODY)
                elif it["k"] == "chart":
                    cd = CategoryChartData(); cd.categories = it["cats"]
                    for name, vals in it["series"]: cd.add_series(name, vals)
                    ct = XL_CHART_TYPE.BAR_CLUSTERED if it["kind"] == "bar" else XL_CHART_TYPE.COLUMN_CLUSTERED
                    ch = sl.shapes.add_chart(ct, Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]), cd).chart
                    ch.font.size = Pt(it["size"]); ch.font.name = BODY_FONT; ch.font.color.rgb = hexc(MUTED)
                    ch.has_title = bool(it["title"])
                    if it["title"]:
                        ch.chart_title.text_frame.text = it["title"]
                        tr = ch.chart_title.text_frame.paragraphs[0].runs[0].font; tr.size = Pt(it["size"] + 1); tr.bold = True; tr.color.rgb = hexc(BODY)
                    ch.has_legend = len(it["series"]) > 1
                    if ch.has_legend: ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
                    plot = ch.plots[0]; plot.gap_width = 60; plot.overlap = -10
                    plot.has_data_labels = True; dl = plot.data_labels
                    dl.number_format = it["fmt"]; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.OUTSIDE_END
                    dl.font.size = Pt(it["size"]); dl.font.color.rgb = hexc(BODY)
                    for ser, col in zip(plot.series, it["colors"]):
                        ser.format.fill.solid(); ser.format.fill.fore_color.rgb = hexc(col)
                    va = ch.value_axis; va.has_major_gridlines = True
                    va.major_gridlines.format.line.color.rgb = hexc(LINE); va.format.line.fill.background()
                    va.tick_labels.font.size = Pt(it["size"] - 1)
                    if it["unit"]: va.has_title = True; va.axis_title.text_frame.text = it["unit"]; va.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(it["size"] - 1)
                    if it["log"]:
                        sc = va._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
                        lb = etree.SubElement(sc, "{http://schemas.openxmlformats.org/drawingml/2006/chart}logBase"); lb.set("val", "10")
                        sc.remove(lb); sc.insert(0, lb)
                    ca = ch.category_axis; ca.format.line.color.rgb = hexc(LINE); ca.tick_labels.font.size = Pt(it["size"])
                    if it["kind"] == "bar": ca.reverse_order = True
                elif it["k"] == "text":
                    tb = sl.shapes.add_textbox(Inches(it["x"]), Inches(it["y"]), Inches(it["w"]), Inches(it["h"]))
                    tf = tb.text_frame; tf.word_wrap = True
                    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[it["anchor"]]
                    for i, p in enumerate(it.get("paras", [])):
                        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                        par.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[it["align"]]
                        par.line_spacing = it["spacing"]; par.space_after = Pt(p["space_after"])
                        run = par.add_run(); run.text = p["txt"]
                        f = run.font; f.name = it["font"]; f.size = Pt(p["size"]); f.bold = p["bold"]; f.italic = p["italic"]
                        f.color.rgb = hexc(p["color"])
            if sp.get("notes"): sl.notes_slide.notes_text_frame.text = sp["notes"]
        prs.save(out)
        return Path(out)

    def save_preview(self, out, title="deck preview"):
        """The same geometry as HTML, so a browser can show what the deck looks like."""
        px = 96
        o = [f"<!doctype html><meta charset='utf-8'><title>{htmlmod.escape(title)}</title><style>",
             "body{margin:0;background:#33383f;font-family:Calibri,Carlito,sans-serif}",
             f".s{{position:relative;width:{W * px}px;height:{H * px}px;margin:18px auto;overflow:hidden}}",
             ".i{position:absolute}", ".t{white-space:pre-wrap}", "</style>"]
        for n, sp in enumerate(self.slides, 1):
            o.append(f"<div class='s' style='background:#{sp['bg']}' id='s{n}'>")
            for it in sp["items"]:
                st = f"left:{it['x'] * px}px;top:{it['y'] * px}px;width:{it['w'] * px}px;height:{it['h'] * px}px"
                if it["k"] == "box":
                    o.append(f"<div class='i' style='{st};background:#{it['fill']};border:{it['lw']}px solid #{it['line']};"
                             f"border-radius:{it['radius'] * px}px;box-sizing:border-box'></div>")
                elif it["k"] == "circle":
                    o.append(f"<div class='i' style='{st};background:#{it['fill']};border-radius:50%'></div>")
                elif it["k"] == "img":
                    o.append(f"<img class='i' style='{st}' src='{it['src']}'>")
                elif it["k"] == "line":
                    o.append(f"<svg class='i' style='left:0;top:0;width:{W * px}px;height:{H * px}px;overflow:visible'><line x1='{it['x1'] * px}' y1='{it['y1'] * px}' "
                             f"x2='{it['x2'] * px}' y2='{it['y2'] * px}' stroke='#{it['color']}' stroke-width='{it['lw'] * 1.33}'/></svg>")
                elif it["k"] == "table":
                    cw = it["col_w"] or [1] * len(it["rows"][0])
                    trs = "".join("<tr>" + "".join(
                        f"<td style='width:{c / sum(cw) * 100:.1f}%;height:{it['row_h'] * px}px;padding:0 0.1in;font-size:{it['size']}pt;"
                        f"background:#{(it['head_fill'] if i == 0 else ('FFFFFF' if i % 2 else 'EEF2F7'))};color:#{'FFFFFF' if i == 0 else BODY};"
                        f"font-weight:{700 if i == 0 or (j == 0 and it['bold_first']) else 400}'>{htmlmod.escape(str(v))}</td>"
                        for j, (v, c) in enumerate(zip(r, cw))) + "</tr>" for i, r in enumerate(it["rows"]))
                    o.append(f"<table class='i' style='{st};border-collapse:collapse;font-family:Calibri,Carlito,sans-serif'>{trs}</table>")
                elif it["k"] == "chart":
                    o.append(f"<div class='i' style='{st}'>{_chart_svg(it, px)}</div>")
                elif it["k"] == "text":
                    al = {"l": "left", "c": "center"}[it["align"]]
                    fam = "Cambria, 'Caladea', serif" if it["font"] == HEAD_FONT else ("'Courier New', 'Liberation Mono', monospace" if it["font"] == MONO_FONT else "Calibri, Carlito, sans-serif")
                    inner = "".join(
                        f"<div style='font-size:{p['size']}pt;color:#{p['color']};font-weight:{700 if p['bold'] else 400};"
                        f"font-style:{'italic' if p['italic'] else 'normal'};line-height:{it['spacing']};"
                        f"margin-bottom:{p['space_after']}pt'>{htmlmod.escape(p['txt'])}</div>" for p in it.get("paras", []))
                    o.append(f"<div class='i t' style='{st};text-align:{al};font-family:{fam};display:flex;"
                             f"flex-direction:column;justify-content:{'center' if it['anchor'] == 'm' else 'flex-start'}'>{inner}</div>")
            o.append(f"<div style='position:absolute;right:10px;bottom:6px;font-size:9pt;color:#9aa4b2'>{n}</div></div>")
            if sp.get("notes"): o.append(f"<div style='width:{W * px}px;margin:-10px auto 30px;color:#cfd6df;font:11pt Calibri,sans-serif'><b>Notes:</b> {htmlmod.escape(sp['notes'])}</div>")
        Path(out).write_text("\n".join(o), encoding="utf-8")
        return Path(out)


def _chart_svg(it, px):
    """An approximation of the native chart for the preview: bars, labels, log or linear axis."""
    import math
    W_, H_ = it["w"] * px, it["h"] * px
    vals = [v for _, vs in it["series"] for v in vs]
    lo = min(vals) / 2 if it["log"] else 0
    hi = max(vals) * (2 if it["log"] else 1.15)
    sc = (lambda v: (math.log10(v) - math.log10(lo)) / (math.log10(hi) - math.log10(lo))) if it["log"] else (lambda v: v / hi)
    top = 40 if (it["title"] or len(it["series"]) > 1) else 10
    out = [f"<svg width='{W_}' height='{H_}' style='font-family:Calibri,sans-serif;font-size:{it['size']}pt'>"]
    if it["title"]: out.append(f"<text x='{W_ / 2}' y='16' text-anchor='middle' font-weight='700' fill='#{BODY}'>{htmlmod.escape(it['title'])}</text>")
    if len(it["series"]) > 1:
        for k, ((nm, _), c) in enumerate(zip(it["series"], it["colors"])):
            out.append(f"<rect x='{W_ / 2 - 80 + k * 110}' y='{top - 16}' width='10' height='10' fill='#{c}'/><text x='{W_ / 2 - 66 + k * 110}' y='{top - 7}' fill='#{MUTED}'>{htmlmod.escape(nm)}</text>")
    ns, nc = len(it["series"]), len(it["cats"])
    if it["kind"] == "bar":
        left, right, bottom = 190, 60, 20
        band = (H_ - top - bottom) / nc
        for ci, cat in enumerate(it["cats"]):
            out.append(f"<text x='{left - 8}' y='{top + band * ci + band / 2 + 4}' text-anchor='end' fill='#{BODY}'>{htmlmod.escape(cat)}</text>")
            for si, ((nm, vs), c) in enumerate(zip(it["series"], it["colors"])):
                bh = band * 0.7 / ns; y = top + band * ci + band * 0.15 + si * bh; w = sc(vs[ci]) * (W_ - left - right)
                out.append(f"<rect x='{left}' y='{y}' width='{max(w, 1)}' height='{bh - 2}' fill='#{c}'/><text x='{left + w + 4}' y='{y + bh / 2 + 4}' fill='#{BODY}'>{vs[ci]:g}</text>")
    else:
        left, bottom = 50, 40
        band = (W_ - left - 10) / nc
        for ci, cat in enumerate(it["cats"]):
            out.append(f"<text x='{left + band * ci + band / 2}' y='{H_ - 18}' text-anchor='middle' fill='#{BODY}'>{htmlmod.escape(cat)}</text>")
            for si, ((nm, vs), c) in enumerate(zip(it["series"], it["colors"])):
                bw = band * 0.7 / ns; x = left + band * ci + band * 0.15 + si * bw; h = sc(vs[ci]) * (H_ - top - bottom)
                out.append(f"<rect x='{x}' y='{H_ - bottom - h}' width='{bw - 2}' height='{max(h, 1)}' fill='#{c}'/><text x='{x + bw / 2}' y='{H_ - bottom - h - 4}' text-anchor='middle' fill='#{BODY}'>{vs[ci]:g}</text>")
    out.append("</svg>")
    return "".join(out)
